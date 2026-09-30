from __future__ import annotations

import hashlib
import re
import unicodedata
from collections import defaultdict
from datetime import datetime

from openpyxl import load_workbook
from sqlalchemy import select

from config.settings import DOWNLOAD_DIR
from config.sources import SINESP_YEAR_XLSX
from database.models import DataQuality, Municipality
from repositories.indicator_repository import IndicatorRepository
from repositories.source_repository import touch_source
from utils.requests import ExternalServiceError, http


def _norm(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(c for c in text if not unicodedata.combining(c)).upper()
    return re.sub(r"[^A-Z0-9]+", "_", text).strip("_")


def _find_column(columns, *terms):
    """Prefer an exact header before considering partial matches."""
    normalized = [(column, _norm(column)) for column in columns]
    wanted_terms = [_norm(term) for term in terms]
    for wanted in wanted_terms:
        for original, clean in normalized:
            if wanted == clean:
                return original
    for wanted in wanted_terms:
        for original, clean in normalized:
            if wanted in clean:
                return original
    return None


def _has_value(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def _number(value: object) -> float | None:
    if not _has_value(value):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _row_measure(row, indexes: dict[str, int], total_col: str | None, victims_col: str | None):
    """Return one non-overlapping measure from a Sinesp row."""
    total = _number(row[indexes[total_col]]) if total_col else None
    if total is not None:
        return total, "ocorrencias"
    victims = _number(row[indexes[victims_col]]) if victims_col else None
    if victims is not None:
        return victims, "vitimas"
    return None, None


def _indicator_code(event: str, scope: str, measure: str) -> str:
    identity = f"{_norm(event)}_{_norm(scope)}_{measure.upper()}".lower()
    digest = hashlib.sha1(identity.encode("utf-8")).hexdigest()[:8]
    return f"crime_{identity[:84]}_{digest}"


class SinespService:
    def __init__(self, session):
        self.session = session
        self.source = touch_source(session, "sinesp")
        self.indicators = IndicatorRepository(session)

    def _municipality(self, codigo_ibge: str) -> Municipality:
        municipality = self.session.get(Municipality, codigo_ibge)
        if municipality is None:
            raise ExternalServiceError(
                "Município não encontrado no cadastro territorial oficial do IBGE."
            )

        same_name_codes = {
            item.codigo_ibge
            for item in self.session.scalars(
                select(Municipality).where(Municipality.uf == municipality.uf)
            )
            if _norm(item.name) == _norm(municipality.name)
        }
        if same_name_codes != {codigo_ibge}:
            raise ExternalServiceError(
                "O nome do município não possui correspondência territorial única dentro da UF."
            )
        return municipality

    def _set_quality(self, codigo_ibge: str, year: int, status: str, coverage: str, notes: str) -> None:
        quality = self.session.scalar(
            select(DataQuality).where(
                DataQuality.dataset == "sinesp",
                DataQuality.codigo_ibge == codigo_ibge,
                DataQuality.year == year,
            )
        )
        if quality is None:
            self.session.add(
                DataQuality(
                    dataset="sinesp",
                    codigo_ibge=codigo_ibge,
                    year=year,
                    status=status,
                    coverage=coverage,
                    notes=notes,
                )
            )
        else:
            quality.status = status
            quality.coverage = coverage
            quality.notes = notes
            quality.checked_at = datetime.utcnow()

    def sync_municipality(self, codigo_ibge: str, year: int) -> int:
        if year < 2015 or year > datetime.now().year:
            raise ExternalServiceError("A série municipal Sinesp VDE publicada cobre anos a partir de 2015.")

        municipality = self._municipality(codigo_ibge)
        source_url = SINESP_YEAR_XLSX.format(year=year)
        destination = DOWNLOAD_DIR / f"sinesp_vde_{year}.xlsx"
        if not destination.exists():
            http.stream_to_file(source_url, destination)

        workbook = load_workbook(destination, read_only=True, data_only=True)
        aggregates: dict[tuple[str, str, str], float] = defaultdict(float)
        matched_municipality = False
        used_name_crosswalk = False
        try:
            for worksheet in workbook.worksheets:
                rows = worksheet.iter_rows(values_only=True)
                raw_headers = next(rows, None)
                if not raw_headers:
                    continue
                headers = [str(value or "") for value in raw_headers]
                indexes = {header: idx for idx, header in enumerate(headers)}

                code_col = _find_column(headers, "codigo ibge", "cod ibge")
                uf_col = _find_column(headers, "uf", "sigla uf")
                municipality_col = _find_column(headers, "municipio", "nome municipio")
                event_col = _find_column(headers, "evento", "indicador", "natureza", "tipo")
                total_col = _find_column(headers, "total", "quantidade", "ocorrencias")
                victims_col = _find_column(headers, "total vitima", "total de vitimas", "vitimas")
                scope_col = _find_column(headers, "abrangencia")
                breakdown_cols = [
                    column
                    for column in (
                        _find_column(headers, "agente"),
                        _find_column(headers, "arma"),
                        _find_column(headers, "faixa etaria"),
                    )
                    if column is not None
                ]

                has_official_code = code_col is not None
                can_crosswalk = uf_col is not None and municipality_col is not None
                if event_col is None or (total_col is None and victims_col is None):
                    continue
                if not has_official_code and not can_crosswalk:
                    continue

                for row in rows:
                    if has_official_code:
                        raw = row[indexes[code_col]]
                        if isinstance(raw, (int, float)):
                            raw_code = str(int(raw)).zfill(7)
                        else:
                            raw_code = re.sub(r"\D", "", str(raw or "")).zfill(7)
                        is_target = raw_code == codigo_ibge
                    else:
                        is_target = (
                            _norm(row[indexes[uf_col]]) == _norm(municipality.uf)
                            and _norm(row[indexes[municipality_col]]) == _norm(municipality.name)
                        )
                        used_name_crosswalk = used_name_crosswalk or is_target
                    if not is_target:
                        continue

                    matched_municipality = True
                    if any(_has_value(row[indexes[column]]) for column in breakdown_cols):
                        continue
                    event = row[indexes[event_col]]
                    if not _has_value(event):
                        continue
                    value, measure = _row_measure(row, indexes, total_col, victims_col)
                    if value is None or measure is None:
                        continue
                    scope = (
                        str(row[indexes[scope_col]]).strip()
                        if scope_col and _has_value(row[indexes[scope_col]])
                        else "Abrangência não informada"
                    )
                    aggregates[(str(event).strip(), scope, measure)] += value
        finally:
            workbook.close()

        if not aggregates:
            reason = (
                "A publicação municipal do Sinesp não contém registros agregáveis para o município/ano."
                if matched_municipality
                else "A publicação municipal do Sinesp não contém o município/ano ou mudou de layout."
            )
            self._set_quality(codigo_ibge, year, "not_available", "Sem registros importáveis", reason)
            raise ExternalServiceError(reason)

        measure_labels = {"ocorrencias": "ocorrências", "vitimas": "vítimas"}
        for (event, scope, measure), value in aggregates.items():
            code = _indicator_code(event, scope, measure)
            label = measure_labels[measure]
            self.indicators.upsert(
                {
                    "codigo_ibge": codigo_ibge,
                    "indicator_code": code,
                    "indicator_name": f"{event} — {label} ({scope})"[:220],
                    "category": "security",
                    "year": year,
                    "value": value,
                    "unit": label.capitalize(),
                    "source_id": self.source.id,
                    "reference_period": str(year),
                    "methodology_version": "Sinesp VDE; total anual por evento, abrangência e tipo de medida",
                    "quality_status": "partial",
                    "original_reference": source_url,
                    "source_record_id": f"SINESP:{codigo_ibge}:{year}:{code}",
                    "is_comparable": True,
                }
            )

        if used_name_crosswalk:
            coverage = "UF + nome oficial IBGE, correspondência única"
            notes = (
                "A fonte não publica código IBGE. O município foi resolvido por igualdade exata após "
                "normalização do nome oficial e da UF; correspondências ausentes ou ambíguas são recusadas."
            )
            status = "partial"
        else:
            coverage = "Código IBGE publicado pela fonte"
            notes = "Associação territorial feita diretamente pelo código IBGE da publicação."
            status = "partial"
        self._set_quality(codigo_ibge, year, status, coverage, notes)
        return len(aggregates)
